import datetime as dt
import json
from typing import Dict, List

from openpyxl import load_workbook

UNIFIED_FIELDS = {
    "a_party",
    "b_party",
    "call_type",
    "direction",
    "start_time",
    "end_time",
    "duration_sec",
    "cell_id",
    "lac_id",
    "imei",
    "imsi",
    "lat",
    "lng",
    "location",
}


OPERATOR_MAPPINGS = {
    "zong": {
        "header_row": 6,
        "columns": {
            "CALL_TYPE": "call_type",
            "MSISDN": "a_party",
            "STRT_TM": "start_time",
            "BNUMBER": "b_party",
            "MINS": "duration_min",
            "SECS": "duration_sec",
            "LAC_ID": "lac_id",
            "CELL_ID": "cell_id",
            "IMEI": "imei",
            "SITE_ADDRESS": "location",
            "LNG": "lng",
            "LAT": "lat",
        },
    },
    "jazz": {
        "header_row": 1,
        "columns": {
            "CallType": "call_type",
            "Aparty": "a_party",
            "BParty": "b_party",
            "Datetime": "start_time",
            "Duration": "duration_sec",
            "cellid": "cell_id",
            "Imsi": "imsi",
            "Imei": "imei",
            "SiteLocation": "location",
        },
    },
    "warid": {
        "header_row": 1,
        "columns": {
            "CallType": "call_type",
            "Aparty": "a_party",
            "BParty": "b_party",
            "Datetime": "start_time",
            "Duration": "duration_sec",
            "cellid": "cell_id",
            "Imsi": "imsi",
            "Imei": "imei",
            "SiteLocation": "location",
        },
    },
    "mobilink": {
        "header_row": 1,
        "columns": {
            "CallType": "call_type",
            "Aparty": "a_party",
            "BParty": "b_party",
            "Datetime": "start_time",
            "Duration": "duration_sec",
            "cellid": "cell_id",
            "Imsi": "imsi",
            "Imei": "imei",
            "SiteLocation": "location",
        },
    },
    "telenor": {
        "header_row": 1,
        "columns": {
            "MSISDN": "a_party",
            "call_org_num": "a_party",
            "CALL_DIALED_NUM": "b_party",
            "IMSI": "imsi",
            "IMEI": "imei",
            "CALL_START_DT_TM": "start_time",
            "CALL_END_DT_TM": "end_time",
            "CALL TYPE": "call_type",
            "DUR": "duration_sec",
            "Lac_Id": "lac_id",
            "Site_Id": "cell_id",
            "Cell_SITE_ID": "cell_id",
            "lat": "lat",
            "longitude": "lng",
            "location": "location",
        },
    },
    "ufone": {
        "header_row": 1,
        "columns": {
            "IMEI": "imei",
            "IMSI": "imsi",
            "A Number": "a_party",
            "B Number": "b_party",
            "Start Time": "start_time",
            "End Time": "end_time",
            "Service Provider": "call_type",
            "Type": "call_type",
            "Direction": "direction",
            "Location": "location",
            "Cell Id": "cell_id",
            "Cell Sector": "lac_id",
            "Latitude": "lat",
            "Longitude": "lng",
            "Duration": "duration_sec",
        },
    },
}


class CDRParsingError(Exception):
    pass


def _normalize_duration(row: Dict) -> int:
    if row.get("duration_sec") is not None:
        try:
            return int(float(row["duration_sec"]))
        except (ValueError, TypeError):
            return 0
    if row.get("duration_min") is not None:
        try:
            minutes = float(row["duration_min"])
        except (ValueError, TypeError):
            minutes = 0
        seconds = 0
        try:
            seconds = float(row.get("duration_sec") or 0)
        except (ValueError, TypeError):
            seconds = 0
        return int(minutes * 60 + seconds)
    return 0


def _parse_datetime(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min).isoformat()
    text = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d-%m-%Y %H:%M:%S",
        "%d-%m-%Y %H:%M",
    ):
        try:
            return dt.datetime.strptime(text, fmt).isoformat()
        except ValueError:
            continue
    return text


def parse_cdr(file_path: str, operator: str) -> List[Dict]:
    operator_key = operator.lower().strip()
    if operator_key not in OPERATOR_MAPPINGS:
        raise CDRParsingError(f"Unsupported operator: {operator}")
    mapping = OPERATOR_MAPPINGS[operator_key]

    workbook = load_workbook(file_path, data_only=True)
    records: List[Dict] = []

    for sheet in workbook.worksheets:
        header_row = mapping["header_row"]
        headers = {}
        for col in range(1, sheet.max_column + 1):
            value = sheet.cell(row=header_row, column=col).value
            if value is None:
                continue
            headers[col] = str(value).strip()
        if not headers:
            continue

        for row_index in range(header_row + 1, sheet.max_row + 1):
            row_data: Dict[str, object] = {}
            empty_row = True
            for col_index, header in headers.items():
                cell_value = sheet.cell(row=row_index, column=col_index).value
                if cell_value not in (None, ""):
                    empty_row = False
                row_data[header] = cell_value
            if empty_row:
                continue

            normalized = {field: None for field in UNIFIED_FIELDS}
            for source, target in mapping["columns"].items():
                normalized[target] = row_data.get(source)

            normalized["duration_sec"] = _normalize_duration(normalized)
            normalized["start_time"] = _parse_datetime(normalized.get("start_time"))
            normalized["end_time"] = _parse_datetime(normalized.get("end_time"))

            record = {
                "a_party": normalized.get("a_party"),
                "b_party": normalized.get("b_party"),
                "call_type": normalized.get("call_type"),
                "direction": normalized.get("direction"),
                "start_time": normalized.get("start_time"),
                "end_time": normalized.get("end_time"),
                "duration_sec": normalized.get("duration_sec"),
                "cell_id": normalized.get("cell_id"),
                "lac_id": normalized.get("lac_id"),
                "imei": normalized.get("imei"),
                "imsi": normalized.get("imsi"),
                "lat": normalized.get("lat"),
                "lng": normalized.get("lng"),
                "location": normalized.get("location"),
                "raw_json": json.dumps(row_data, default=str),
            }
            records.append(record)

    if not records:
        raise CDRParsingError("No CDR rows found in the uploaded file.")
    return records
